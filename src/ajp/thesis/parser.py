"""T16: PDF (pdfplumber primary, pypdf fallback) and DOCX (python-docx) -> one neutral ``ThesisDoc`` model,
shared by the formatting checks (checks.py) and the similarity pipeline (similarity/). Every field either
comes directly from the document or is left ``None``/empty - nothing here guesses a value it cannot support
(D6's rule extends past ad-listing extraction to this module too).

DOCX gives real page-layout metadata (margins, fonts, line spacing) via python-docx's section/style objects.
PDF gives real per-page text and word-level font metadata via pdfplumber, from which headings are inferred by
relative font size/boldness - a heuristic, never claimed to be as reliable as DOCX's structured styles, and
PDF ``page_style`` values stay ``None`` unless directly observable (pdfplumber exposes page width/height, not
the document's original margin settings, since those are baked into the fixed page geometry once rendered to
PDF - so PDF page_style.source is always 'pdf_heuristic' and margins/font/line-spacing stay unset there).
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pypdf
from docx import Document as DocxDocument
from docx.enum.text import WD_LINE_SPACING
from docx.opc.exceptions import PackageNotFoundError
from pdfplumber import PDF as PdfPlumberPDF

from ajp.thesis.sr_script import to_latin_for_matching

EMU_PER_CM = 360000


class ParseError(ValueError):
    pass


# ---------------------------------------------------------------------------------------------- model -------
@dataclass
class Section:
    title: str
    level: int                                # 1 = top-level chapter/heading, 2 = subsection, ...
    numbering: Optional[str] = None           # "2.1" etc - None when the heading has no leading number
    page_start: Optional[int] = None          # 1-based; PDF only (DOCX has no fixed pagination pre-print)
    raw_heading_text: str = ""


@dataclass
class Abstract:
    text: str
    lang: Optional[str] = None                # 'sr' | 'en' | None (undetermined)
    script: Optional[str] = None              # 'latin' | 'cyrillic' | 'mixed' | None
    page_start: Optional[int] = None


@dataclass
class TocEntry:
    title: str
    page: Optional[int] = None
    numbering: Optional[str] = None


@dataclass
class FigureOrTable:
    kind: str                                 # 'figure' | 'table'
    caption: str
    number: Optional[str] = None
    page: Optional[int] = None


@dataclass
class Reference:
    raw: str
    index: Optional[int] = None               # the [1]/1. marker in the reference LIST itself, if present


@dataclass
class PageStyle:
    source: str                               # 'docx' | 'pdf_heuristic'
    margins_cm: Optional[Dict[str, float]] = None
    font_name: Optional[str] = None
    font_size_pt: Optional[float] = None
    line_spacing: Optional[float] = None      # multiple of single spacing; None if not a simple multiple
    justification: Optional[str] = None       # 'left' | 'right' | 'center' | 'justify' | None


@dataclass
class ParagraphRecord:
    style_name: str                           # the DOCX style name verbatim, e.g. 'Heading 1', 'Normal'
    text: str


@dataclass
class ThesisDoc:
    source_format: str                        # 'docx' | 'pdf'
    page_count: Optional[int]
    full_text: str
    page_texts: List[str] = field(default_factory=list)   # PDF only; empty for DOCX
    paragraphs: List[ParagraphRecord] = field(default_factory=list)  # DOCX only, original order+style, verbatim
    # - export/docx_export.py reconstructs the document from this (never re-typing content), and re-applies
    # the template's own font/size/spacing/margins on top of it. Empty for PDF: there is no comparable
    # reconstructible paragraph stream once a document has been rendered to fixed PDF layout.
    metadata: Dict[str, Any] = field(default_factory=dict)
    sections: List[Section] = field(default_factory=list)
    abstracts: List[Abstract] = field(default_factory=list)
    toc: List[TocEntry] = field(default_factory=list)
    figures: List[FigureOrTable] = field(default_factory=list)
    tables: List[FigureOrTable] = field(default_factory=list)
    references: List[Reference] = field(default_factory=list)
    appendices: List[Section] = field(default_factory=list)
    page_style: Optional[PageStyle] = None
    has_page_numbers: Optional[bool] = None   # None = undetermined (e.g. DOCX without a rendered page count)


# --------------------------------------------------------------------------------------- shared helpers -----
_CYRILLIC_RE = re.compile(r"[Ѐ-ӿ]")
_LATIN_LETTER_RE = re.compile(r"[A-Za-zĐđŠšČčĆćŽž]")
_NUMBERING_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(\S.*)$")
_SR_ABSTRACT_MARK = re.compile(r"\b(apstrakt|sa[žz]etak|klju[cč]ne re[cč]i)\b", re.IGNORECASE)
_EN_ABSTRACT_MARK = re.compile(r"\b(abstract|keywords)\b", re.IGNORECASE)
#  Latin and Cyrillic literal alternatives side by side (not transliterated at match time) so a captured
#  group - a figure caption's free text, for instance - comes back in the document's own original script,
#  never silently converted. Landmarks with no captured free text (abstract/toc/references headings) could
#  use transliteration safely, but keeping every landmark on the same pattern is simpler to audit.
_ABSTRACT_HEADING = re.compile(r"^\s*(apstrakt|sa[žz]etak|abstract|апстракт|сажетак)\s*$", re.IGNORECASE)
_TOC_HEADING = re.compile(r"^\s*(sadr[žz]aj|table of contents|contents|садржај)\s*$", re.IGNORECASE)
_TOC_LINE = re.compile(r"^(.{3,120}?)[\s.․…]{2,}\s*(\d{1,4})\s*$")
_REFERENCES_HEADING = re.compile(r"^\s*(literatura|reference[ns]?|bibliografija|литература|референце|библиографија)\s*$",
                                 re.IGNORECASE)
_FIGURE_CAPTION = re.compile(r"^\s*(slika|figure|fig\.|слика)\s*(\d+(?:\.\d+)*)\s*[:.]\s*(\S.*)$", re.IGNORECASE)
_TABLE_CAPTION = re.compile(r"^\s*(tabela|table|табела)\s*(\d+(?:\.\d+)*)\s*[:.]\s*(\S.*)$", re.IGNORECASE)
_HEADING_KEYWORDS = {
    "uvod", "pregled literature", "pregled literature/stanja u oblasti", "ciljevi i hipoteze istraživanja",
    "materijal i metode", "metodologija", "rezultati", "diskusija", "zaključak", "literatura", "prilozi",
    "biografija", "sadržaj", "abstract", "introduction", "conclusion",
}


def detect_script(text: str) -> Optional[str]:
    """'latin' | 'cyrillic' | 'mixed' | None (no letters at all)."""
    cyr = len(_CYRILLIC_RE.findall(text))
    lat = len(_LATIN_LETTER_RE.findall(text))
    if cyr == 0 and lat == 0:
        return None
    if cyr > 0 and lat > 0:
        # a handful of stray Latin characters (units, abbreviations) inside an otherwise-Cyrillic block is
        # still "cyrillic" for the reader's purposes; only call it genuinely mixed past a low noise floor.
        minor = min(cyr, lat) / max(cyr, lat)
        return "mixed" if minor > 0.08 else ("cyrillic" if cyr > lat else "latin")
    return "cyrillic" if cyr > 0 else "latin"


def detect_abstract_lang(text: str) -> Optional[str]:
    if _SR_ABSTRACT_MARK.search(text):
        return "sr"
    if _EN_ABSTRACT_MARK.search(text):
        return "en"
    script = detect_script(text)
    if script == "cyrillic":
        return "sr"
    return None


def _parse_numbering(line: str) -> "tuple[Optional[str], str]":
    m = _NUMBERING_RE.match(line.strip())
    if m:
        return m.group(1), m.group(2).strip()
    return None, line.strip()


def _looks_like_heading(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 120:
        return False
    numbering, rest = _parse_numbering(stripped)
    if numbering:
        return True
    return to_latin_for_matching(stripped).lower().rstrip(":") in _HEADING_KEYWORDS


# ---------------------------------------------------------------------------------------------- DOCX --------
def _docx_page_style(doc: DocxDocument) -> PageStyle:
    sec = doc.sections[0] if doc.sections else None
    margins = None
    if sec is not None:
        def cm(v: Optional[int]) -> Optional[float]:
            return round(v / EMU_PER_CM, 2) if v is not None else None
        margins = {"top": cm(sec.top_margin), "bottom": cm(sec.bottom_margin),
                   "left": cm(sec.left_margin), "right": cm(sec.right_margin)}
    normal = doc.styles["Normal"] if "Normal" in [s.name for s in doc.styles] else None
    font_name = normal.font.name if normal is not None else None
    font_size = normal.font.size.pt if normal is not None and normal.font.size is not None else None
    line_spacing, justification = None, None
    if normal is not None and normal.paragraph_format is not None:
        pf = normal.paragraph_format
        if pf.line_spacing_rule == WD_LINE_SPACING.SINGLE:
            line_spacing = 1.0
        elif pf.line_spacing_rule == WD_LINE_SPACING.DOUBLE:
            line_spacing = 2.0
        elif pf.line_spacing_rule == WD_LINE_SPACING.ONE_POINT_FIVE:
            line_spacing = 1.5
        elif isinstance(pf.line_spacing, (int, float)):
            line_spacing = round(float(pf.line_spacing), 2)
        if pf.alignment is not None:
            justification = pf.alignment.name.lower()          # WD_ALIGN_PARAGRAPH.JUSTIFY -> 'justify'
    return PageStyle(source="docx", margins_cm=margins, font_name=font_name, font_size_pt=font_size,
                     line_spacing=line_spacing, justification=justification)


def parse_docx(path: Union[str, Path]) -> ThesisDoc:
    try:
        doc = DocxDocument(str(path))
    except PackageNotFoundError as e:
        raise ParseError(f"not a valid .docx file: {e}") from e

    paragraph_records = [ParagraphRecord(style_name=(p.style.name if p.style is not None else "Normal") or "Normal",
                                         text=p.text) for p in doc.paragraphs]
    full_text = "\n".join(p.text for p in paragraph_records)

    sections: List[Section] = []
    abstracts: List[Abstract] = []
    toc: List[TocEntry] = []
    figures: List[FigureOrTable] = []
    tables: List[FigureOrTable] = []
    references: List[Reference] = []

    mode: Optional[str] = None                # None | 'abstract' | 'toc' | 'references'
    abstract_buf: List[str] = []
    ref_index = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        style_name = (p.style.name or "").lower() if p.style is not None else ""
        is_heading_style = style_name.startswith("heading") or style_name.startswith("naslov")

        if _ABSTRACT_HEADING.match(text):
            if mode == "abstract" and abstract_buf:
                abstracts.append(_finish_abstract(abstract_buf))
            mode, abstract_buf = "abstract", []
            continue
        if _TOC_HEADING.match(text):
            if mode == "abstract" and abstract_buf:
                abstracts.append(_finish_abstract(abstract_buf))
                abstract_buf = []
            # the TOC heading itself is still a real top-level heading (required_sections checks for it) -
            # only the mode switch, not this section record, is conditioned on anything else.
            sections.append(Section(title=text, level=1, numbering=None, raw_heading_text=text))
            mode = "toc"
            continue
        if _REFERENCES_HEADING.match(text) or (is_heading_style and text.strip().lower() in {"literatura", "references"}):
            if mode == "abstract" and abstract_buf:
                abstracts.append(_finish_abstract(abstract_buf))
                abstract_buf = []
            sections.append(Section(title=text, level=1, numbering=None, raw_heading_text=text))
            mode = "references"
            continue
        if is_heading_style and text:
            # a real new heading (Prilozi, Biografija, ...) always ends whatever mode we were in, including
            # 'references' - otherwise every heading after the reference list gets swallowed as a fake
            # reference entry (found live: 'Prilozi'/'Biografija' + their body text counted as 4 extra refs).
            if mode == "abstract" and abstract_buf:
                abstracts.append(_finish_abstract(abstract_buf))
                abstract_buf = []
            mode = None

        if mode == "abstract":
            if text:
                abstract_buf.append(text)
            continue
        if mode == "toc":
            m = _TOC_LINE.match(text)
            if m:
                numbering, title = _parse_numbering(m.group(1))
                toc.append(TocEntry(title=title, page=int(m.group(2)), numbering=numbering))
            continue
        if mode == "references":
            if text:
                ref_index += 1
                references.append(Reference(raw=text, index=ref_index))
            continue

        if is_heading_style and text:
            level = 1
            digits = re.search(r"(\d+)", style_name)
            if digits:
                level = int(digits.group(1))
            numbering, title = _parse_numbering(text)
            sections.append(Section(title=title, level=level, numbering=numbering, raw_heading_text=text))
            continue

        fm = _FIGURE_CAPTION.match(text)
        if fm:
            figures.append(FigureOrTable(kind="figure", number=fm.group(2), caption=fm.group(3)))
            continue
        tm = _TABLE_CAPTION.match(text)
        if tm:
            tables.append(FigureOrTable(kind="table", number=tm.group(2), caption=tm.group(3)))
            continue

    if mode == "abstract" and abstract_buf:
        abstracts.append(_finish_abstract(abstract_buf))

    return ThesisDoc(
        source_format="docx", page_count=None, full_text=full_text, page_texts=[], paragraphs=paragraph_records,
        metadata=_docx_metadata(doc), sections=sections, abstracts=abstracts, toc=toc,
        figures=figures, tables=tables, references=references, appendices=[],
        page_style=_docx_page_style(doc), has_page_numbers=None,
    )


def _finish_abstract(lines: List[str]) -> Abstract:
    text = "\n".join(lines).strip()
    return Abstract(text=text, lang=detect_abstract_lang(text), script=detect_script(text))


def _docx_metadata(doc: DocxDocument) -> Dict[str, Any]:
    cp = doc.core_properties
    return {"title": cp.title or None, "author": cp.author or None, "subject": cp.subject or None,
            "created": cp.created.isoformat() if cp.created else None}


# ---------------------------------------------------------------------------------------------- PDF ---------
def parse_pdf(path: Union[str, Path]) -> ThesisDoc:
    try:
        return _parse_pdf_pdfplumber(path)
    except Exception:                                        # noqa: BLE001 - fall back to pypdf, text-only
        return _parse_pdf_pypdf(path)


def _parse_pdf_pdfplumber(path: Union[str, Path]) -> ThesisDoc:
    with PdfPlumberPDF.open(str(path)) as pdf:
        page_texts = [(pg.extract_text() or "") for pg in pdf.pages]
        page_count = len(pdf.pages)
        full_text = "\n".join(page_texts)

        sections: List[Section] = []
        abstracts: List[Abstract] = []
        toc: List[TocEntry] = []
        figures: List[FigureOrTable] = []
        tables: List[FigureOrTable] = []
        references: List[Reference] = []
        mode: Optional[str] = None
        abstract_buf: List[str] = []
        ref_index = 0

        body_size = _median_font_size(pdf.pages)

        for page_no, pg in enumerate(pdf.pages, start=1):
            words = pg.extract_words(extra_attrs=["size", "fontname"]) or []
            lines = _group_words_into_lines(words)
            for line_text, avg_size, is_bold in lines:
                text = line_text.strip()
                if not text:
                    continue
                if _ABSTRACT_HEADING.match(text):
                    if mode == "abstract" and abstract_buf:
                        a = _finish_abstract(abstract_buf)
                        a.page_start = a.page_start
                        abstracts.append(a)
                    mode, abstract_buf = "abstract", []
                    continue
                if _TOC_HEADING.match(text):
                    if mode == "abstract" and abstract_buf:
                        abstracts.append(_finish_abstract(abstract_buf))
                        abstract_buf = []
                    sections.append(Section(title=text, level=1, numbering=None, page_start=page_no, raw_heading_text=text))
                    mode = "toc"
                    continue
                if _REFERENCES_HEADING.match(text):
                    if mode == "abstract" and abstract_buf:
                        abstracts.append(_finish_abstract(abstract_buf))
                        abstract_buf = []
                    sections.append(Section(title=text, level=1, numbering=None, page_start=page_no, raw_heading_text=text))
                    mode = "references"
                    continue

                if mode == "abstract":
                    if _looks_like_heading(text) and text.strip().lower() not in _HEADING_KEYWORDS:
                        pass  # short heading-shaped lines inside an abstract are still abstract prose
                    if avg_size > body_size * 1.15 and _looks_like_heading(text):
                        mode = None
                    else:
                        abstract_buf.append(text)
                        continue
                if mode == "toc":
                    m = _TOC_LINE.match(text)
                    if m:
                        numbering, title = _parse_numbering(m.group(1))
                        toc.append(TocEntry(title=title, page=int(m.group(2)), numbering=numbering))
                        continue
                    if _looks_like_heading(text) and avg_size > body_size * 1.15:
                        mode = None
                    else:
                        continue
                if mode == "references":
                    # same fix as the DOCX path: a real new heading (Prilozi, Biografija, ...) after the
                    # reference list must end references mode, not be swallowed as one more reference entry.
                    if avg_size > body_size * 1.15 and _looks_like_heading(text):
                        mode = None
                    else:
                        ref_index += 1
                        references.append(Reference(raw=text, index=ref_index))
                        continue

                fm = _FIGURE_CAPTION.match(text)
                if fm:
                    figures.append(FigureOrTable(kind="figure", number=fm.group(2), caption=fm.group(3), page=page_no))
                    continue
                tm = _TABLE_CAPTION.match(text)
                if tm:
                    tables.append(FigureOrTable(kind="table", number=tm.group(2), caption=tm.group(3), page=page_no))
                    continue

                if (avg_size > body_size * 1.15 or is_bold) and _looks_like_heading(text):
                    numbering, title = _parse_numbering(text)
                    level = 1 if numbering is None or "." not in numbering else numbering.count(".") + 1
                    sections.append(Section(title=title, level=level, numbering=numbering,
                                            page_start=page_no, raw_heading_text=text))

        if mode == "abstract" and abstract_buf:
            abstracts.append(_finish_abstract(abstract_buf))

        has_page_numbers = _pdf_has_page_numbers(pdf.pages)
        return ThesisDoc(
            source_format="pdf", page_count=page_count, full_text=full_text, page_texts=page_texts,
            metadata=_pdf_metadata(pdf), sections=sections, abstracts=abstracts, toc=toc,
            figures=figures, tables=tables, references=references, appendices=[],
            page_style=PageStyle(source="pdf_heuristic"), has_page_numbers=has_page_numbers,
        )


def _pdf_metadata(pdf: PdfPlumberPDF) -> Dict[str, Any]:
    meta = pdf.metadata or {}
    return {"title": meta.get("Title") or None, "author": meta.get("Author") or None,
            "subject": meta.get("Subject") or None, "created": meta.get("CreationDate") or None}


def _median_font_size(pages) -> float:
    sizes: List[float] = []
    for pg in pages[: min(len(pages), 10)]:
        for w in pg.extract_words(extra_attrs=["size"]) or []:
            if "size" in w:
                sizes.append(w["size"])
    if not sizes:
        return 10.0
    sizes.sort()
    return sizes[len(sizes) // 2]


def _group_words_into_lines(words: List[Dict[str, Any]]):
    """pdfplumber's extract_words gives one dict per word; group consecutive words on ~the same baseline
    ('top') into a line, and report that line's average font size and whether most of it is bold."""
    lines: List[List[Dict[str, Any]]] = []
    for w in words:
        if lines and abs(lines[-1][-1]["top"] - w["top"]) < 3:
            lines[-1].append(w)
        else:
            lines.append([w])
    out = []
    for line in lines:
        text = " ".join(w["text"] for w in line)
        sizes = [w.get("size", 0) for w in line]
        avg_size = sum(sizes) / len(sizes) if sizes else 0
        bold_count = sum(1 for w in line if "bold" in w.get("fontname", "").lower())
        out.append((text, avg_size, bold_count > len(line) / 2))
    return out


def _pdf_has_page_numbers(pages) -> bool:
    """A short line consisting only of a number, near the top or bottom 10% of the page, on most pages."""
    hits = 0
    checked = pages[1: min(len(pages), 11)]                   # skip the title page
    for pg in checked:
        h = float(pg.height)
        for w in pg.extract_words() or []:
            if re.fullmatch(r"\d{1,4}", w["text"]) and (w["top"] < h * 0.1 or w["bottom"] > h * 0.9):
                hits += 1
                break
    return bool(checked) and hits >= max(1, len(checked) // 2)


def _parse_pdf_pypdf(path: Union[str, Path]) -> ThesisDoc:
    """Fallback when pdfplumber cannot open the file: text-only, no layout heuristics (D6 - unsupported
    structure is left empty, never guessed from a format this route cannot see)."""
    reader = pypdf.PdfReader(str(path))
    page_texts = [(pg.extract_text() or "") for pg in reader.pages]
    full_text = "\n".join(page_texts)
    meta = reader.metadata or {}
    return ThesisDoc(
        source_format="pdf", page_count=len(reader.pages), full_text=full_text, page_texts=page_texts,
        metadata={"title": getattr(meta, "title", None), "author": getattr(meta, "author", None)},
        page_style=PageStyle(source="pdf_heuristic"),
    )


# ---------------------------------------------------------------------------------------------- entry -------
def parse_document(path: Union[str, Path]) -> ThesisDoc:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".docx":
        return parse_docx(p)
    if suffix == ".pdf":
        return parse_pdf(p)
    raise ParseError(f"unsupported file type: {suffix!r} (only .docx and .pdf are supported)")


def nfkc(text: str) -> str:
    return unicodedata.normalize("NFKC", text)
